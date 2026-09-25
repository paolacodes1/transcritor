"""Turn Whisper segments into .txt / .srt / .docx files."""
import datetime as dt

PARAGRAPH_GAP = 2.0  # seconds of silence that starts a new paragraph
PARAGRAPH_MAX_CHARS = 900

LANG_NAMES = {"pt": "Português", "en": "English", "es": "Español", "fr": "Français",
              "it": "Italiano", "de": "Deutsch"}


def ts(seconds, srt=False):
    seconds = max(0.0, seconds)
    h = int(seconds // 3600)
    m = int(seconds % 3600 // 60)
    s = int(seconds % 60)
    if srt:
        ms = int(round((seconds - int(seconds)) * 1000))
        if ms == 1000:
            ms = 999
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def paragraphs(segments):
    """Group segments into readable paragraphs: [(start_time, text)]."""
    paras, cur, cur_start, last_end = [], [], None, None
    for seg in segments:
        new_para = (
            cur
            and (
                seg["start"] - last_end >= PARAGRAPH_GAP
                or (sum(len(t) for t in cur) > PARAGRAPH_MAX_CHARS and cur[-1][-1:] in ".?!")
            )
        )
        if new_para:
            paras.append((cur_start, " ".join(cur)))
            cur, cur_start = [], None
        if cur_start is None:
            cur_start = seg["start"]
        cur.append(seg["text"])
        last_end = seg["end"]
    if cur:
        paras.append((cur_start, " ".join(cur)))
    return paras


def to_txt(segments):
    return "\n\n".join(p for _, p in paragraphs(segments)) + "\n"


def to_timestamped_txt(segments):
    return "\n\n".join(f"[{ts(t)}] {p}" for t, p in paragraphs(segments)) + "\n"


def to_srt(segments):
    out = []
    for i, s in enumerate(segments, 1):
        out.append(f"{i}\n{ts(s['start'], True)} --> {ts(s['end'], True)}\n{s['text']}\n")
    return "\n".join(out)


def to_docx(segments, path, title, language, duration):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(title, level=1)
    meta = doc.add_paragraph()
    run = meta.add_run(
        f"{dt.date.today().strftime('%d/%m/%Y')}  ·  "
        f"{LANG_NAMES.get(language, language or '')}  ·  {ts(duration)}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for start, text in paragraphs(segments):
        p = doc.add_paragraph()
        stamp = p.add_run(f"[{ts(start)}]  ")
        stamp.font.size = Pt(9)
        stamp.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.add_run(text)
        p.paragraph_format.space_after = Pt(8)
    doc.save(path)
